import openpyxl
from openpyxl.styles import Font, PatternFill
from pathlib import Path
from typing import List, Dict, Any, Optional
from .config import OPTIMIZED_CALENDAR_PATH

class CalendarManager:
    """Manages reading and updating Zara_Khan_Optimized_Calendar.xlsx."""

    def __init__(self, calendar_path: Optional[Path] = None):
        self.calendar_path = calendar_path or OPTIMIZED_CALENDAR_PATH
        if not self.calendar_path.exists():
            raise FileNotFoundError(f"Optimized calendar not found at: {self.calendar_path}")

    def load_workbook(self):
        return openpyxl.load_workbook(self.calendar_path)

    def get_headers(self, ws) -> Dict[str, int]:
        """Returns a mapping of header name to 1-based column index."""
        headers = {}
        for col in range(1, ws.max_column + 1):
            val = ws.cell(1, col).value
            if val:
                headers[str(val).strip()] = col
        return headers

    def get_assets(
        self,
        day: Optional[str] = None,
        asset_id: Optional[str] = None,
        image_only: bool = True,
        pending_only: bool = False,
        limit: Optional[int] = None
    ) -> List[Dict[str, Any]]:
        """Retrieves asset records matching filters."""
        wb = self.load_workbook()
        ws = wb.active
        headers = self.get_headers(ws)

        records = []
        for r in range(2, ws.max_row + 1):
            aid = ws.cell(r, headers.get("Asset ID", 1)).value
            if not aid:
                continue

            day_val = ws.cell(r, headers.get("Day", 2)).value
            post_type = ws.cell(r, headers.get("Post Type", 3)).value
            asset_type = ws.cell(r, headers.get("Asset Type", 4)).value
            is_image = bool(ws.cell(r, headers.get("Is Image Asset", 5)).value)
            status = ws.cell(r, headers.get("Generation Status", 6)).value
            skip_reason = ws.cell(r, headers.get("Skip Reason", 7)).value
            desc = ws.cell(r, headers.get("Content Description", 8)).value
            notes = ws.cell(r, headers.get("Notes (Overrides)", 9)).value
            caption = ws.cell(r, headers.get("Caption", 10)).value
            music = ws.cell(r, headers.get("Music", 11)).value
            visual_url = ws.cell(r, headers.get("Visual Reference URL", 12)).value
            reel_url = ws.cell(r, headers.get("Reel Reference URL", 13)).value
            ref_path = ws.cell(r, headers.get("Reference Media Path", 14)).value
            final_prompt = ws.cell(r, headers.get("Final LLM Prompt", 15)).value
            neg_prompt = ws.cell(r, headers.get("Negative Prompt", 16)).value
            out_path = ws.cell(r, headers.get("Output Path", 17)).value
            source_row = ws.cell(r, headers.get("Source Row", 20)).value

            # Filter conditions
            if image_only and not is_image:
                continue
            if pending_only and status != "PENDING":
                continue
            if day:
                norm_day = str(day).strip().lower()
                row_day = str(day_val or "").strip().lower()
                # Exact match or normalized day match
                target_day_str = norm_day if norm_day.startswith("day") else f"day {norm_day}"
                if row_day != target_day_str:
                    continue
            if asset_id:
                if str(asset_id).strip().upper() != str(aid).strip().upper():
                    continue

            record = {
                "row_index": r,
                "asset_id": str(aid).strip(),
                "day": str(day_val or "").strip(),
                "post_type": str(post_type or "").strip(),
                "asset_type": str(asset_type or "").strip(),
                "is_image_asset": is_image,
                "generation_status": str(status or "").strip(),
                "skip_reason": skip_reason,
                "content_description": desc,
                "notes": notes,
                "caption": caption,
                "music": music,
                "visual_reference_url": visual_url,
                "reel_reference_url": reel_url,
                "reference_media_path": ref_path,
                "final_llm_prompt": final_prompt,
                "negative_prompt": neg_prompt,
                "output_path": out_path,
                "source_row": source_row
            }
            records.append(record)

            if limit and len(records) >= limit:
                break

        wb.close()
        return records

    def update_asset(
        self,
        asset_id: str,
        status: Optional[str] = None,
        llm_prompt: Optional[str] = None,
        negative_prompt: Optional[str] = None,
        ref_media_path: Optional[str] = None,
        output_path: Optional[str] = None,
        skip_reason: Optional[str] = None
    ) -> bool:
        """Updates fields for a specific asset ID and saves the workbook."""
        wb = self.load_workbook()
        ws = wb.active
        headers = self.get_headers(ws)

        target_row = None
        aid_col = headers.get("Asset ID", 1)
        for r in range(2, ws.max_row + 1):
            val = ws.cell(r, aid_col).value
            if val and str(val).strip().upper() == asset_id.strip().upper():
                target_row = r
                break

        if not target_row:
            wb.close()
            return False

        # Status fills
        fills = {
            "GENERATED": PatternFill(start_color="D9EAD3", end_color="D9EAD3", fill_type="solid"), # soft green
            "DRY_RUN": PatternFill(start_color="D0E0E3", end_color="D0E0E3", fill_type="solid"),   # soft cyan
            "FAILED": PatternFill(start_color="FCE5CD", end_color="FCE5CD", fill_type="solid"),    # soft orange/red
            "PENDING": PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid"),   # soft yellow
            "SKIPPED": PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")    # gray
        }

        if status:
            cell = ws.cell(target_row, headers["Generation Status"], status)
            if status in fills:
                cell.fill = fills[status]
                cell.font = Font(name="Calibri", size=10, bold=True)

        if llm_prompt is not None:
            ws.cell(target_row, headers["Final LLM Prompt"], llm_prompt)
        if negative_prompt is not None:
            ws.cell(target_row, headers["Negative Prompt"], negative_prompt)
        if ref_media_path is not None:
            ws.cell(target_row, headers["Reference Media Path"], ref_media_path)
        if output_path is not None:
            ws.cell(target_row, headers["Output Path"], output_path)
        if skip_reason is not None:
            ws.cell(target_row, headers["Skip Reason"], skip_reason)

        wb.save(self.calendar_path)
        wb.close()
        return True

    def get_summary(self) -> Dict[str, int]:
        """Returns statistics of the pipeline calendar."""
        wb = self.load_workbook()
        ws = wb.active
        headers = self.get_headers(ws)

        stats = {
            "total_records": 0,
            "image_assets": 0,
            "skipped_assets": 0,
            "pending": 0,
            "generated": 0,
            "dry_run": 0,
            "failed": 0
        }

        for r in range(2, ws.max_row + 1):
            aid = ws.cell(r, headers.get("Asset ID", 1)).value
            if not aid:
                continue
            stats["total_records"] += 1
            is_image = bool(ws.cell(r, headers.get("Is Image Asset", 5)).value)
            status = str(ws.cell(r, headers.get("Generation Status", 6)).value or "").upper()

            if is_image:
                stats["image_assets"] += 1
            else:
                stats["skipped_assets"] += 1

            if status == "PENDING":
                stats["pending"] += 1
            elif status == "GENERATED":
                stats["generated"] += 1
            elif status == "DRY_RUN":
                stats["dry_run"] += 1
            elif status == "FAILED":
                stats["failed"] += 1

        wb.close()
        return stats
